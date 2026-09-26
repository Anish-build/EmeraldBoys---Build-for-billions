import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# Folder containing your CSV files
CSV_FOLDER = Path("jalsetu_dummy_dataset/jalsetu_dummy_dataset")

# Excel file to create
OUTPUT_FILE = "jalsetu_dummy_dataset_export.xlsx"

# Create Excel workbook
workbook = Workbook()

# Remove default empty sheet
workbook.remove(workbook.active)

# Find every CSV file
csv_files = sorted(CSV_FOLDER.glob("*.csv"))

if not csv_files:
    print("ERROR: No CSV files found!")
    print(f"Checked folder: {CSV_FOLDER}")
    exit()

# Process every CSV file
for csv_file in csv_files:

    # Sheet name = CSV filename without .csv
    sheet_name = csv_file.stem[:31]

    print(f"Exporting: {csv_file.name}")

    # Create sheet
    sheet = workbook.create_sheet(title=sheet_name)

    # Read CSV
    with open(csv_file, "r", encoding="utf-8-sig", newline="") as file:
        reader = csv.reader(file)
        rows = list(reader)

    if not rows:
        continue

    # Write everything into Excel
    for row_number, row in enumerate(rows, start=1):
        for column_number, value in enumerate(row, start=1):
            cell = sheet.cell(
                row=row_number,
                column=column_number,
                value=value
            )

            # Make first row bold
            if row_number == 1:
                cell.font = Font(bold=True)

    # Freeze header row
    sheet.freeze_panes = "A2"

    # Add filter to header
    if sheet.max_column > 0 and sheet.max_row > 1:
        sheet.auto_filter.ref = sheet.dimensions

    # Automatically adjust column widths
    for column in sheet.columns:

        max_length = 0

        for cell in column:
            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        column_letter = get_column_letter(
            column[0].column
        )

        # Limit width so extremely long text doesn't make the sheet huge
        sheet.column_dimensions[column_letter].width = min(
            max_length + 2,
            40
        )

# Save Excel workbook
workbook.save(OUTPUT_FILE)

print()
print("=" * 60)
print("CSV EXPORT COMPLETE")
print("=" * 60)
print(f"Created: {OUTPUT_FILE}")
print()
print("CSV files exported:")

for csv_file in csv_files:
    print(f"  ✓ {csv_file.name}")

print("=" * 60)