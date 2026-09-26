import sqlite3
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

# Database to export
DB_NAME = "jalsetu_import_test.db"

# Output Excel file
OUTPUT_FILE = "jalsetu_database_export.xlsx"

# Tables to export
TABLES = [
    "pumps",
    "mechanics",
    "inspections",
    "diagnoses",
    "repairs",
    "verifications",
    "audit_log"
]

# Connect to database
conn = sqlite3.connect(DB_NAME)
cursor = conn.cursor()

# Create Excel workbook
workbook = Workbook()

# Remove default sheet
default_sheet = workbook.active
workbook.remove(default_sheet)

# Export every table
for table in TABLES:

    # Create worksheet
    sheet = workbook.create_sheet(title=table[:31])

    # Get column names
    cursor.execute(f"PRAGMA table_info({table})")
    columns = [column[1] for column in cursor.fetchall()]

    # Write column headers
    for col_num, column_name in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=col_num)
        cell.value = column_name
        cell.font = cell.font.copy(bold=True)

    # Get all data
    cursor.execute(f"SELECT * FROM {table}")
    rows = cursor.fetchall()

    # Write data
    for row_num, row in enumerate(rows, start=2):
        for col_num, value in enumerate(row, start=1):
            sheet.cell(
                row=row_num,
                column=col_num,
                value=value
            )

    # Freeze header row
    sheet.freeze_panes = "A2"

    # Add filter
    if columns:
        sheet.auto_filter.ref = sheet.dimensions

    # Automatically adjust column widths
    for column_cells in sheet.columns:
        max_length = 0

        for cell in column_cells:
            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        column_letter = get_column_letter(
            column_cells[0].column
        )

        sheet.column_dimensions[column_letter].width = min(
            max_length + 2,
            40
        )

# Save Excel file
workbook.save(OUTPUT_FILE)

# Close database
conn.close()

print()
print("=" * 60)
print("DATABASE EXPORT COMPLETE")
print("=" * 60)
print(f"Database : {DB_NAME}")
print(f"Excel    : {OUTPUT_FILE}")
print()
print("Exported tables:")

for table in TABLES:
    print(f"  ✓ {table}")

print("=" * 60)